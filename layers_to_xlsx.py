#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
layers_to_xlsx.py —— 把逐层统计 CSV 转成带图表的 Excel

为什么用 openpyxl 而不是 matplotlib
    openpyxl 生成的是 Excel 原生图表对象：数据源与图表同在文件内，
    点开单元格能看到底层数据，图表样式也能随时调整。
    比导出的静态 PNG 更实用 —— 一份文件同时承载数据与可视化，便于归档与二次编辑。

用法
    python layers_to_xlsx.py sample_layers.csv out.xlsx --title "Benchy 0.2mm 层高"
"""
import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def read_layers(csv_path):
    """读 CSV。用最朴素的方式解析，避免依赖 pandas。"""
    rows = []
    text = Path(csv_path).read_text(encoding="utf-8-sig")
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        return rows

    header = [h.strip() for h in lines[0].split(",")]
    for line in lines[1:]:
        parts = [p.strip() for p in line.split(",")]
        if len(parts) != len(header):
            continue
        record = {}
        for key, value in zip(header, parts):
            # 数值列转成 float，转不了就保留字符串
            try:
                record[key] = float(value)
            except ValueError:
                record[key] = value
        rows.append(record)
    return rows


def build_workbook(rows, title):
    """创建 Excel：左侧放数据，右侧放两张图。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "逐层数据"

    # ---------- 标题行 ----------
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A1:D1")

    # ---------- 表头 ----------
    headers = ["层号", "Z高度(mm)", "耗时(min)", "挤出(mm)"]
    keys = ["layer", "z_mm", "time_min", "extrude_mm"]
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    for col, (label, key) in enumerate(zip(headers, keys), start=1):
        cell = ws.cell(row=2, column=col, value=label)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")

    # ---------- 数据行 ----------
    for i, record in enumerate(rows, start=3):
        for col, key in enumerate(keys, start=1):
            value = record.get(key, "")
            cell = ws.cell(row=i, column=col, value=value)
            if key != "layer":
                cell.number_format = "0.00"

    first_data_row = 3
    last_data_row = 2 + len(rows)
    if not rows:
        return wb, None, None

    # ---------- 图表 1：Z 高度 vs 逐层耗时 ----------
    chart1 = LineChart()
    chart1.title = "逐层耗时 (Time per layer)"
    chart1.y_axis.title = "Z高度 (mm)"
    chart1.x_axis.title = "耗时 (min)"
    chart1.height = 8
    chart1.width = 12
    data = Reference(ws, min_col=2, min_row=2, max_row=last_data_row)
    cats = Reference(ws, min_col=3, min_row=first_data_row, max_row=last_data_row)
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    # 图表放在 F 列，避免压住数据
    ws.add_chart(chart1, "F2")

    # ---------- 图表 2：Z 高度 vs 逐层挤出量 ----------
    chart2 = LineChart()
    chart2.title = "逐层挤出量 (Extrusion per layer)"
    chart2.y_axis.title = "Z高度 (mm)"
    chart2.x_axis.title = "挤出 (mm)"
    chart2.height = 8
    chart2.width = 12
    data2 = Reference(ws, min_col=2, min_row=2, max_row=last_data_row)
    cats2 = Reference(ws, min_col=4, min_row=first_data_row, max_row=last_data_row)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    ws.add_chart(chart2, "F20")

    # ---------- 列宽 ----------
    for col, width in zip("ABCD", (10, 12, 12, 12)):
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A3"

    # ---------- 第二个工作表：汇总 ----------
    summary = wb.create_sheet("汇总")
    summary["A1"] = "汇总指标"
    summary["A1"].font = Font(bold=True, size=13)

    total_time = sum(float(r.get("time_min", 0) or 0) for r in rows)
    total_ext = sum(float(r.get("extrude_mm", 0) or 0) for r in rows)
    max_time = max((float(r.get("time_min", 0) or 0) for r in rows), default=0)

    metrics = [
        ("总层数", len(rows)),
        ("总耗时(min)", round(total_time, 2)),
        ("总挤出(mm)", round(total_ext, 1)),
        ("单层最大耗时(min)", round(max_time, 2)),
        ("平均单层耗时(min)", round(total_time / len(rows), 3) if rows else 0),
    ]
    for i, (label, value) in enumerate(metrics, start=3):
        summary.cell(row=i, column=1, value=label).font = Font(bold=True)
        summary.cell(row=i, column=2, value=value)

    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 14

    return wb, chart1, chart2


def main(argv=None):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

    parser = argparse.ArgumentParser(description="逐层统计 CSV -> 带图表的 Excel")
    parser.add_argument("csv", help="layers CSV 文件（gcode_analyzer.py --csv 产出）")
    parser.add_argument("out", help="输出 xlsx 路径")
    parser.add_argument("--title", default="G代码逐层分析", help="表格标题")
    args = parser.parse_args(argv)

    rows = read_layers(args.csv)
    if not rows:
        print(f"[错误] 没读到数据: {args.csv}", file=sys.stderr)
        return 1

    wb, _, _ = build_workbook(rows, args.title)
    wb.save(args.out)
    print(f"[已导出] {args.out}  ({len(rows)} 层)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
